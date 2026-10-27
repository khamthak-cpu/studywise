"""Validated upload storage and extraction for conversation attachments."""

from __future__ import annotations

import hashlib
import mimetypes
import uuid
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

import pymupdf as fitz
from docx import Document
from fastapi import HTTPException, UploadFile

from app.services.productivity import connect


# Project paths and limits
PROJECT_ROOT = Path(__file__).resolve().parents[3]
UPLOAD_DIR = PROJECT_ROOT / "data" / "uploads"

MAX_FILE_BYTES = 15 * 1024 * 1024  # 15 MB
MAX_EXTRACTED_CHARS = 120_000

ALLOWED_EXTENSIONS = {
    ".pdf",
    ".png",
    ".jpg",
    ".jpeg",
    ".webp",
    ".docx",
    ".txt",
}

IMAGE_EXTENSIONS = {".png", ".jpg", ".jpeg", ".webp"}


def _now() -> str:
    """Return the current UTC timestamp."""
    return datetime.now(timezone.utc).isoformat()


def init_attachments_db() -> None:
    """Create the attachment table and index if they do not exist."""
    with connect() as db:
        db.executescript(
            """
            CREATE TABLE IF NOT EXISTS attachments (
                id TEXT PRIMARY KEY,
                student_id TEXT NOT NULL,
                conversation_id INTEGER NOT NULL,
                filename TEXT NOT NULL,
                extension TEXT NOT NULL,
                content_type TEXT NOT NULL,
                storage_path TEXT NOT NULL,
                extracted_text TEXT NOT NULL DEFAULT '',
                sha256 TEXT NOT NULL,
                created_at TEXT NOT NULL,
                FOREIGN KEY(conversation_id)
                    REFERENCES conversations(id) ON DELETE CASCADE
            );

            CREATE INDEX IF NOT EXISTS idx_attachments_conversation
            ON attachments(student_id, conversation_id, created_at);
            """
        )


def _extract_text(path: Path, extension: str) -> str:
    """Extract searchable text from supported document formats."""

    if extension == ".pdf":
        chunks: list[str] = []

        with fitz.open(path) as pdf:
            for page_number, page in enumerate(pdf, start=1):
                text = page.get_text("text").strip()

                if text:
                    chunks.append(
                        f"[Page {page_number}]\n{text}"
                    )

        return "\n\n".join(chunks)[:MAX_EXTRACTED_CHARS]

    if extension == ".docx":
        document = Document(path)

        chunks = [
            paragraph.text.strip()
            for paragraph in document.paragraphs
            if paragraph.text.strip()
        ]

        for table in document.tables:
            for row in table.rows:
                chunks.append(
                    " | ".join(cell.text.strip() for cell in row.cells)
                )

        return "\n".join(chunks)[:MAX_EXTRACTED_CHARS]

    if extension == ".txt":
        return path.read_text(
            encoding="utf-8",
            errors="replace",
        )[:MAX_EXTRACTED_CHARS]

    # Images are processed separately by the vision model.
    return ""


def _is_valid_file(path: Path, extension: str) -> bool:
    """Check basic file signatures rather than trusting filenames alone."""

    try:
        with path.open("rb") as file:
            header = file.read(16)

        if extension == ".pdf":
            return header.startswith(b"%PDF-")

        if extension == ".png":
            return header.startswith(b"\x89PNG\r\n\x1a\n")

        if extension in {".jpg", ".jpeg"}:
            return header.startswith(b"\xff\xd8\xff")

        if extension == ".webp":
            return (
                len(header) >= 12
                and header[:4] == b"RIFF"
                and header[8:12] == b"WEBP"
            )

        if extension == ".docx":
            # DOCX files are ZIP containers.
            return header.startswith(b"PK")

        if extension == ".txt":
            return b"\x00" not in header

    except OSError:
        return False

    return False


def get_conversation_attachments(
    student_id: str,
    conversation_id: int,
) -> list[dict[str, Any]]:
    """Return attachments belonging to this student and conversation."""

    init_attachments_db()

    with connect() as db:
        rows = db.execute(
            """
            SELECT
                id,
                student_id,
                conversation_id,
                filename,
                extension,
                content_type,
                storage_path,
                extracted_text,
                sha256,
                created_at
            FROM attachments
            WHERE student_id = ? AND conversation_id = ?
            ORDER BY created_at ASC
            """,
            (student_id, conversation_id),
        ).fetchall()

    return [dict(row) for row in rows]


def get_selected_attachments(
    student_id: str,
    conversation_id: int,
    ids: list[str],
) -> list[dict[str, Any]]:
    """Return requested attachments, or the five most recent by default."""

    all_items = get_conversation_attachments(
        student_id,
        conversation_id,
    )

    if not ids:
        return all_items[-5:]

    requested_ids = set(ids)

    selected = [
        item
        for item in all_items
        if item["id"] in requested_ids
    ]

    if len(selected) != len(requested_ids):
        raise HTTPException(
            status_code=404,
            detail=(
                "One or more attachments were not found "
                "in this conversation."
            ),
        )

    return selected


async def save_upload(
    file: UploadFile,
    student_id: str,
    conversation_id: int,
) -> dict[str, Any]:
    """Validate, store, extract, and register an uploaded attachment."""

    original_name = file.filename or "upload"
    filename = Path(original_name.replace("\\", "/")).name[:255]
    extension = Path(filename).suffix.lower()

    if extension not in ALLOWED_EXTENSIONS:
        await file.close()
        raise HTTPException(
            status_code=415,
            detail=(
                "Supported files: PDF, PNG, JPG, WEBP, DOCX and TXT."
            ),
        )

    init_attachments_db()

    # Verify that the conversation exists and belongs to this student.
    with connect() as db:
        owner = db.execute(
            """
            SELECT id
            FROM conversations
            WHERE id = ? AND student_id = ?
            """,
            (conversation_id, student_id),
        ).fetchone()

    if owner is None:
        await file.close()
        raise HTTPException(
            status_code=404,
            detail="Conversation not found for this student.",
        )

    UPLOAD_DIR.mkdir(parents=True, exist_ok=True)

    attachment_id = uuid.uuid4().hex
    path = (UPLOAD_DIR / f"{attachment_id}{extension}").resolve()

    if UPLOAD_DIR.resolve() not in path.parents:
        await file.close()
        raise HTTPException(
            status_code=400,
            detail="Invalid upload path.",
        )

    digest = hashlib.sha256()
    size = 0

    try:
        # Stream to disk and enforce the size limit while reading.
        with path.open("xb") as destination:
            while True:
                chunk = await file.read(1024 * 1024)

                if not chunk:
                    break

                size += len(chunk)

                if size > MAX_FILE_BYTES:
                    raise HTTPException(
                        status_code=413,
                        detail="Each file must be 15 MB or smaller.",
                    )

                digest.update(chunk)
                destination.write(chunk)

        if size == 0:
            raise HTTPException(
                status_code=400,
                detail="The uploaded file is empty.",
            )

        if not _is_valid_file(path, extension):
            raise HTTPException(
                status_code=415,
                detail=(
                    "The file content does not match its extension "
                    "or the file is invalid."
                ),
            )

        try:
            extracted_text = _extract_text(path, extension)
        except Exception as exc:
            raise HTTPException(
                status_code=422,
                detail=(
                    "Could not process this document. "
                    "Please check that it is valid and not corrupted."
                ),
            ) from exc

        content_type = (
            mimetypes.guess_type(filename)[0]
            or "application/octet-stream"
        )

        with connect() as db:
            db.execute(
                """
                INSERT INTO attachments (
                    id,
                    student_id,
                    conversation_id,
                    filename,
                    extension,
                    content_type,
                    storage_path,
                    extracted_text,
                    sha256,
                    created_at
                )
                VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                """,
                (
                    attachment_id,
                    student_id,
                    conversation_id,
                    filename,
                    extension,
                    content_type,
                    str(path),
                    extracted_text,
                    digest.hexdigest(),
                    _now(),
                ),
            )

        return {
            "id": attachment_id,
            "filename": filename,
            "content_type": content_type,
            "size": size,
            "kind": (
                "image"
                if extension in IMAGE_EXTENSIONS
                else "document"
            ),
            "text_extracted": bool(extracted_text),
        }

    except Exception:
        # Remove incomplete files if validation or database insertion fails.
        path.unlink(missing_ok=True)
        raise

    finally:
        await file.close()
